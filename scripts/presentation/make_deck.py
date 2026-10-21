#!/usr/bin/env python
"""Generate docs/ONEZENTRA_WORKFLOW.pptx — deck visual (diagram + ikon + screenshot aplikasi)."""
import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT = os.environ.get("DECK_OUT", str(ROOT / "docs" / "ONEZENTRA_WORKFLOW.pptx"))
SH = os.environ.get("SHOTS_DIR", str(ROOT / "storage" / "app" / "public" / "presentation" / "shots"))
LOGO = str(ROOT / "onezentra-web" / "public" / "LOGO.png")

PRIMARY = RGBColor(0x4F, 0x46, 0xE5)
DARK = RGBColor(0x31, 0x2E, 0x81)
ACCENT = RGBColor(0x0D, 0x94, 0x88)
TEXT = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xEE, 0xF2, 0xFF)
CARD = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0xDC, 0x26, 0x26)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
GREEN = RGBColor(0x15, 0x80, 0x3D)
NAVY = RGBColor(0x0F, 0x17, 0x2A)
SOFTRED = RGBColor(0xFE, 0xF2, 0xF2)
SOFTGRN = RGBColor(0xEC, 0xFD, 0xF5)
SOFTYEL = RGBColor(0xFE, 0xF9, 0xC3)

W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
_page = {"n": 0}


def rect(slide, left, top, width, height, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, adj=None):
    sh = slide.shapes.add_shape(shape, int(left), int(top), int(width), int(height))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    return sh


def textbox(slide, left, top, width, height, text="", size=18, bold=False, color=TEXT,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb


def bullets(slide, items, left, top, width, height, size=16, gap=6, color=TEXT, marker=ACCENT):
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in items:
        lvl, txt = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        p.level = min(lvl, 4)
        r0 = p.add_run()
        r0.text = "▸ "
        r0.font.size = Pt(size)
        r0.font.bold = True
        r0.font.color.rgb = marker if lvl == 0 else MUTED
        r0.font.name = "Calibri"
        for i, part in enumerate(txt.split("**")):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size)
            r.font.bold = (i % 2 == 1)
            r.font.color.rgb = color if lvl == 0 else MUTED
            r.font.name = "Calibri"
    return tb


def header(slide, title, kicker=None):
    rect(slide, 0, 0, Inches(0.28), H, fill=PRIMARY)
    top = Inches(0.42)
    if kicker:
        textbox(slide, Inches(0.7), top, Inches(11.6), Inches(0.3), kicker.upper(), size=12, bold=True, color=ACCENT)
        top = Inches(0.75)
    textbox(slide, Inches(0.7), top, Inches(11.9), Inches(0.6), title, size=30, bold=True, color=DARK)
    rect(slide, Inches(0.72), Inches(1.45) if kicker else Inches(1.18), Inches(1.5), Pt(4), fill=ACCENT)


def footer(slide, label):
    _page["n"] += 1
    rect(slide, Inches(0.28), H - Inches(0.42), W - Inches(0.28), Pt(1), fill=BORDER)
    textbox(slide, Inches(0.72), H - Inches(0.36), Inches(9.5), Inches(0.28),
            "ONEZENTRA Warehouse System  ·  " + label, size=10, color=MUTED)
    textbox(slide, Inches(12.2), H - Inches(0.36), Inches(0.7), Inches(0.28),
            str(_page["n"]), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def new(title=None, kicker=None, foot="Workflow, Fungsi & Tujuan"):
    s = prs.slides.add_slide(BLANK)
    if title:
        header(s, title, kicker)
    footer(s, foot)
    return s


def card(slide, left, top, width, height, heading, lines, head_color=PRIMARY, size=13, glyph=None):
    rect(slide, left, top, width, height, fill=CARD, line=BORDER)
    rect(slide, left, top, width, Pt(4), fill=head_color)
    hx = left + Inches(0.16)
    if glyph:
        icon_badge(slide, hx, top + Inches(0.16), Inches(0.34), glyph, head_color, fs=14)
        hx += Inches(0.46)
    textbox(slide, hx, top + Inches(0.17), width - (hx - left) - Inches(0.16), Inches(0.34),
            heading, size=15, bold=True, color=head_color)
    bullets(slide, [(0, l) for l in lines], left + Inches(0.16), top + Inches(0.6),
            width - Inches(0.32), height - Inches(0.72), size=size, gap=5)


def icon_badge(slide, left, top, size, glyph, color=PRIMARY, fs=18):
    rect(slide, left, top, size, size, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.25)
    tb = textbox(slide, left, top, size, size, glyph, size=fs, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return tb


def flow(slide, steps, top, left=Inches(0.7), total_w=Inches(11.9), height=Inches(1.05),
         label_size=14, sub_size=10.5, gap=Inches(0.12)):
    n = len(steps)
    bw = int((total_w - gap * (n - 1)) / n)
    for i, st in enumerate(steps):
        main = st[0]
        sub = st[1] if len(st) > 1 else ""
        glyph = st[2] if len(st) > 2 else None
        color = PRIMARY if i % 2 == 0 else ACCENT
        sh = rect(slide, left + i * (bw + gap), top, bw, height, fill=color,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = (glyph + "  " if glyph else "") + main
        r.font.size = Pt(label_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = "Calibri"
        if sub:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = sub
            r2.font.size = Pt(sub_size)
            r2.font.color.rgb = WHITE
            r2.font.name = "Calibri"
        if i < n - 1:
            rect(slide, left + i * (bw + gap) + bw - Inches(0.02), top + height / 3,
                 gap + Inches(0.06), height / 3, fill=BORDER, shape=MSO_SHAPE.RIGHT_ARROW)


def pill(slide, left, top, width, height, text, color, fs=11, bold=True):
    rect(slide, left, top, width, height, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    textbox(slide, left, top, width, height, text, size=fs, bold=bold, color=WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def shot(slide, name, left, top, width, caption=None, cap_color=MUTED, max_h=None):
    """Insert screenshot dengan bingkai + caption. max_h membatasi tinggi (lebar mengecil otomatis)."""
    from PIL import Image
    path = f"{SH}/{name}.png"
    with Image.open(path) as im:
        ratio = im.height / im.width
    if max_h and int(width * ratio) > int(max_h):
        width = int(max_h / ratio)
    height = int(width * ratio)
    pad = Inches(0.06)
    rect(slide, left - pad, top - pad, width + 2 * pad, height + 2 * pad, fill=WHITE, line=BORDER)
    slide.shapes.add_picture(path, int(left), int(top), width=int(width))
    if caption:
        textbox(slide, left, top + height + Inches(0.12), width, Inches(0.3), caption,
                size=11.5, bold=True, color=cap_color)
    return height


def table(slide, headers, rows, left, top, width, col_ratios=None, size=12, head_size=12,
          row_h=0.32, head_h=0.34):
    nrows, ncols = len(rows) + 1, len(headers)
    gt = slide.shapes.add_table(nrows, ncols, int(left), int(top), int(width),
                                Inches(head_h + row_h * len(rows))).table
    if col_ratios:
        total = sum(col_ratios)
        for i, r in enumerate(col_ratios):
            gt.columns[i].width = Emu(int(width * r / total))
    gt.rows[0].height = Inches(head_h)
    for i in range(1, nrows):
        gt.rows[i].height = Inches(row_h)
    for c, htxt in enumerate(headers):
        cell = gt.cell(0, c)
        cell.text = htxt
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        for r in cell.text_frame.paragraphs[0].runs:
            r.font.size = Pt(head_size)
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.name = "Calibri"
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = gt.cell(ri, ci)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else CARD
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            for r in cell.text_frame.paragraphs[0].runs:
                r.font.size = Pt(size)
                r.font.color.rgb = DARK if ci == 0 else TEXT
                r.font.bold = (ci == 0)
                r.font.name = "Calibri"
    return gt


# ══════════════════════════ 1. COVER
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=DARK)
rect(s, 0, 0, Inches(0.35), H, fill=ACCENT)
rect(s, Inches(7.9), 0, Inches(5.44), H, fill=RGBColor(0x3F, 0x3A, 0xA8))
try:
    from PIL import Image
    with Image.open(LOGO) as im:
        lr = im.height / im.width
    s.shapes.add_picture(LOGO, Inches(0.95), Inches(0.75), width=Inches(2.3))
except Exception:
    pass
textbox(s, Inches(0.95), Inches(1.85), Inches(7.0), Inches(0.4), "SISTEM MANAJEMEN GUDANG",
        size=15, bold=True, color=ACCENT)
textbox(s, Inches(0.95), Inches(2.3), Inches(7.2), Inches(2.0), "ONEZENTRA\nWarehouse System",
        size=46, bold=True, color=WHITE, spacing=1.05)
rect(s, Inches(0.97), Inches(4.35), Inches(1.8), Pt(5), fill=ACCENT)
textbox(s, Inches(0.95), Inches(4.7), Inches(7.0), Inches(0.9),
        "Workflow · Fungsi · Tujuan Aplikasi", size=25, color=WHITE)
textbox(s, Inches(0.95), Inches(6.15), Inches(7.0), Inches(0.7),
        "Next.js 16 (Web)  ·  Laravel 11 API  ·  PostgreSQL/SQLite  ·  Flutter (Mobile)\n"
        "Dokumen pendukung: docs/WORKFLOW_TUJUAN_SISTEM.md · docs/project_flow.md",
        size=12, color=RGBColor(0xC7, 0xD2, 0xFE))
chips = [("🚚", "Alur 4 tahap"), ("📦", "Batch otomatis"), ("🛡️", "RBAC & audit"), ("📈", "Laporan & biaya")]
textbox(s, Inches(8.5), Inches(1.7), Inches(4.4), Inches(0.4), "INTI SISTEM", size=14, bold=True, color=ACCENT)
for i, (g, t) in enumerate(chips):
    top = Inches(2.3 + i * 1.0)
    rect(s, Inches(8.5), top, Inches(4.2), Inches(0.8), fill=RGBColor(0x4C, 0x46, 0xB8),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
    textbox(s, Inches(8.7), top, Inches(0.6), Inches(0.8), g, size=22, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, color=WHITE)
    textbox(s, Inches(9.4), top, Inches(3.2), Inches(0.8), t, size=16, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
_page["n"] += 1

# ══════════════════════════ 2. AGENDA
s = new("Agenda", "Daftar Isi")
agenda = [
    ("1", "Tujuan Aplikasi", "Masalah yang dipecahkan & sasaran"),
    ("2", "Tampilan Sistem", "Dashboard & layar utama"),
    ("3", "Arsitektur & Modul", "Empat lapisan · enam grup menu"),
    ("4", "Workflow Utama", "Replenishment → Gate → Validasi → Check Out"),
    ("5", "Alur Data & Barcode", "Scan sekali, batch & label otomatis"),
    ("6", "Workflow Pendukung & Modul Detail", "Opname, biaya, laporan, master data"),
    ("7", "Keamanan & Peran", "RBAC, audit trail, tanggung jawab"),
    ("8", "Manfaat & Batasan", "KPI dan modul yang belum ada"),
]
for i, (num, head, sub) in enumerate(agenda):
    col, row = i % 2, i // 2
    left = Inches(0.7) + col * Inches(6.15)
    top = Inches(1.75) + row * Inches(1.3)
    rect(s, left, top, Inches(5.8), Inches(1.05), fill=CARD, line=BORDER)
    rect(s, left, top, Inches(0.7), Inches(1.05), fill=PRIMARY if col == 0 else ACCENT)
    textbox(s, left, top + Inches(0.24), Inches(0.7), Inches(0.6), num, size=26, bold=True,
            color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, left + Inches(0.9), top + Inches(0.2), Inches(4.7), Inches(0.4), head, size=17, bold=True, color=DARK)
    textbox(s, left + Inches(0.9), top + Inches(0.58), Inches(4.7), Inches(0.35), sub, size=12.5, color=MUTED)

# ══════════════════════════ 3. TUJUAN
s = new("Tujuan Aplikasi", "Mengapa sistem ini dibuat")
textbox(s, Inches(0.7), Inches(1.68), Inches(5.7), Inches(0.35), "❌  MASALAH PADA PROSES MANUAL", size=13, bold=True, color=WARN)
rect(s, Inches(0.7), Inches(2.08), Inches(5.7), Inches(4.5), fill=SOFTRED, line=RGBColor(0xFE, 0xCA, 0xCA))
bullets(s, [
    "Rencana kirim & Surat Jalan berbasis kertas",
    "Barang dibongkar tanpa pencocokan dokumen",
    "Selisih qty baru ketahuan saat stok opname",
    "Batch & masa berlaku produk tidak terpantau",
    "Biaya disetujui lewat chat — tanpa jejak audit",
    "Hak akses data tidak jelas antar fungsi",
], Inches(0.95), Inches(2.35), Inches(5.2), Inches(4.0), size=15, gap=14, color=TEXT)

textbox(s, Inches(6.85), Inches(1.68), Inches(5.7), Inches(0.35), "✅  TUJUAN SISTEM", size=13, bold=True, color=ACCENT)
rect(s, Inches(6.85), Inches(2.08), Inches(5.75), Inches(4.5), fill=SOFTGRN, line=RGBColor(0x99, 0xF6, 0xE4))
bullets(s, [
    "**Kontrol alur barang** — satu jalur wajib untuk setiap truk",
    "**Ketepatan data** — produk discan sekali, diteruskan otomatis",
    "**Ketertelusuran** — status, batch, aktivitas user tercatat",
    "**Kepatuhan finansial** — biaya cair setelah approval berjenjang",
    "**Keamanan akses** — permission per aksi, sesi cookie",
], Inches(7.1), Inches(2.35), Inches(5.25), Inches(4.0), size=15, gap=17, color=TEXT)

# ══════════════════════════ 4. TAMPILAN SISTEM
s = new("Tampilan Sistem", "Dua layar paling sering dipakai")
shot(s, "dashboard", Inches(0.7), Inches(1.78), Inches(5.9),
     "Dashboard — KPI, perlu tindakan, progres rencana kirim")
shot(s, "npn", Inches(6.85), Inches(1.78), Inches(5.75),
     "Data Master NPN — katalog produk + import/export Excel")

# ══════════════════════════ 5. ARSITEKTUR
s = new("Arsitektur Sistem", "Empat lapisan")
flow(s, [("Frontend Web", "Next.js 16 · React 19", "🖥️"),
         ("Proxy / Auth Guard", "proxy.ts · cookie sesi", "🔐"),
         ("Backend API", "Laravel 11 · /api/v1", "⚙️"),
         ("Database", "PostgreSQL · SQLite", "🗄️")],
     Inches(1.72), height=Inches(1.1))
cards = [
    ("onezentra-web/", "🌐", PRIMARY, [
        "**32 halaman** App Router, Tailwind 4, Zustand",
        "Guard auth via **proxy.ts**",
        "Teruskan /api/* & /sanctum/* ke backend",
        "PDF client-side (Surat Jalan & Biaya)",
        "Cetak label ESC/POS via Web Bluetooth",
    ]),
    ("app/ — Laravel API", "⚙️", ACCENT, [
        "**32 controller** prefix /api/v1",
        "37 model · state machine terkunci",
        "Middleware permission:* + track.activity",
        "4 service: approval, email, opname, resolver",
        "20+ export Excel + template import",
    ]),
    ("onezentra_mobile/", "📱", PRIMARY, [
        "Flutter — login & akses API yang sama",
        "Bearer token khusus mobile",
        "Provider + GoRouter, theme mengikuti web",
        "Builder label ESC/POS siap dipindah native",
    ]),
    ("Konsep Kunci", "🔑", ACCENT, [
        "**Stateless API** — semua data lewat REST",
        "**Scan sekali** di Validasi SJ",
        "**Transisi sah** saja (422 bila ilegal)",
        "**Audit trail** — user_activity_logs",
    ]),
]
for i, (h, g, c, l) in enumerate(cards):
    card(s, Inches(0.7) + i * Inches(3.05), Inches(3.2), Inches(2.85), Inches(3.35), h, l, c, size=12, glyph=g)

# ══════════════════════════ 6. PETA MODUL
s = new("Peta Modul & Fungsi", "Enam grup menu")
groups = [
    ("Operasional", "🚚", PRIMARY, [
        "Replenishment — buat & jadwalkan kirim",
        "Check In Point — gate-in kendaraan",
        "Validasi Dokumen — cocokkan SJ vs fisik",
        "Check Out Point — selesaikan bongkar",
        "Inbound (SJ) — master penerimaan depo",
    ]),
    ("Opname", "🧾", ACCENT, [
        "Batch Expired — monitoring mingguan",
        "Sesi Opname — verifikasi lapangan",
        "AMAN / PERHATIAN / KRITIS / EXPIRED",
    ]),
    ("Master Data", "🗂️", PRIMARY, [
        "NPN — katalog produk",
        "Cluster Principal — pengelompokan",
        "Standar Umur — umur simpan standar",
        "Standarisasi — tier & tumpukan",
    ]),
    ("Biaya", "💰", ACCENT, [
        "Biaya THL — tenaga harian lepas",
        "Biaya Retribusi — biaya per dokumen",
        "Tarif Retribusi — tarif per depo",
    ]),
    ("Laporan", "📊", PRIMARY, [
        "Replenishment · Inbound · Batch Expired",
        "Selisih Qty · Penerimaan & Expired",
        "Biaya THL · Biaya Retribusi",
        "Filter tanggal + depo → ekspor Excel",
    ]),
    ("Sistem", "🛡️", ACCENT, [
        "Users · Roles & Permissions",
        "Pengaturan Approval · Reset Data",
        "Setting Email · Penerima Email",
        "Dashboard KPI & aktivitas",
    ]),
]
for i, (h, g, c, l) in enumerate(groups):
    card(s, Inches(0.7) + (i % 3) * Inches(4.05), Inches(1.7) + (i // 3) * Inches(2.55),
         Inches(3.8), Inches(2.35), h, l, c, size=12.5, glyph=g)

# ══════════════════════════ 7. ALUR BESAR
s = new("Workflow Utama — Alur Besar", "Satu jalur wajib untuk setiap truk")
flow(s, [("1. REPLENISHMENT", "Admin Planner", "📋"),
         ("2. CHECK IN", "Security", "🚧"),
         ("3. VALIDASI SJ", "Cheker", "✅"),
         ("4. CHECK OUT", "Cheker", "📦")],
     Inches(1.68), height=Inches(1.15), label_size=16, sub_size=13)
flow(s, [("DRAFT → SCHEDULED", "rencana dijadwalkan"),
         ("CHECKED_IN", "kendaraan tercatat masuk"),
         ("VERIFIED", "qty, kondisi, expired + batch"),
         ("IN_PROGRESS → COMPLETED", "bongkar selesai")],
     Inches(3.05), height=Inches(1.0), label_size=13, sub_size=11)
bullets(s, [
    "**Satu-satunya titik scan produk** adalah Validasi SJ — hasilnya diteruskan otomatis ke Receiving.",
    "Status hanya berpindah lewat **transisi sah**; transisi ilegal ditolak **HTTP 422**.",
    "Setiap tahap memicu **notifikasi email** dan menyimpan **aktivitas user**.",
], Inches(0.7), Inches(4.4), Inches(11.9), Inches(1.8), size=15.5, gap=13)

# ══════════════════════════ 8. SWIMLANE
s = new("Workflow Utama — Siapa Mengerjakan Apa", "Pembagian tugas antar peran")
lanes = [("Admin Planner", PRIMARY), ("Security", ORANGE), ("Cheker", ACCENT), ("Sistem", DARK)]
lane_top, lane_h, lane_gap = Inches(1.75), Inches(0.92), Inches(0.16)
label_w = Inches(2.0)
cols = [("TAHAP 1", "Replenishment"), ("TAHAP 2", "Check In"), ("TAHAP 3", "Validasi SJ"), ("TAHAP 4", "Check Out")]
col_w = Inches(2.45)
col_left = Inches(2.85)
for j, (h1, h2) in enumerate(cols):
    rect(s, col_left + j * (col_w + Inches(0.15)), lane_top - Inches(0.5), col_w, Inches(0.44), fill=NAVY)
    textbox(s, col_left + j * (col_w + Inches(0.15)), lane_top - Inches(0.5), col_w, Inches(0.44),
            f"{h1} · {h2}", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
matrix = [
    # (lane, col) -> (aktivitas, warna)
    ((0, 0), ("Buat rencana · isi dokumen & barang", PRIMARY)),
    ((0, 1), ("Jadwalkan (SCHEDULED)", PRIMARY)),
    ((1, 1), ("Scan barcode · gate-in", ORANGE)),
    ((2, 2), ("Scan item · qty · expired · kondisi", ACCENT)),
    ((2, 3), ("Scan PO · bongkar · selesai", ACCENT)),
    ((3, 0), ("Email TO BLC, CC Cheker", MUTED)),
    ((3, 2), ("Batch BATCH-/LBL- + Receiving otomatis", MUTED)),
    ((3, 3), ("Dock AVAILABLE · email selesai", MUTED)),
]
lookup = {}
for (r, c), v in matrix:
    lookup[(r, c)] = v
for i, (lname, lcol) in enumerate(lanes):
    top = lane_top + i * (lane_h + lane_gap)
    rect(s, Inches(0.7), top, label_w, lane_h, fill=lcol, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.15)
    textbox(s, Inches(0.7), top, label_w, lane_h, lname, size=13, bold=True, color=WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for j in range(4):
        left = col_left + j * (col_w + Inches(0.15))
        rect(s, left, top, col_w, lane_h, fill=CARD, line=BORDER)
        if (i, j) in lookup:
            txt, col = lookup[(i, j)]
            rect(s, left, top, col_w, lane_h, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
            textbox(s, left + Inches(0.12), top, col_w - Inches(0.24), lane_h, txt, size=12, bold=True,
                    color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            rect(s, left + col_w / 2 - Inches(0.22), top + lane_h / 2 - Pt(1.2), Inches(0.44), Pt(2.4), fill=BORDER)
textbox(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.4),
        "Panah alur: Planner (1) → Security (2) → Cheker (3) → Sistem (auto) → Cheker (4) → Sistem (auto)",
        size=13, bold=True, color=DARK)

# ══════════════════════════ 9. STATE MACHINE
s = new("State Machine Rencana Kirim", "Aturan perpindahan status")
flow(s, [("DRAFT", "✍️"), ("SCHEDULED", "🗓️"), ("CHECKED_IN", "🚧"), ("VERIFIED", "✅"),
         ("IN_PROGRESS", "📦"), ("COMPLETED", "🏁")],
     Inches(1.72), height=Inches(0.95), label_size=13, sub_size=10)
rect(s, Inches(0.7), Inches(3.0), Inches(11.9), Inches(0.72), fill=SOFTRED, line=RGBColor(0xFE, 0xCA, 0xCA))
textbox(s, Inches(0.95), Inches(3.12), Inches(11.4), Inches(0.5),
        "Dari DRAFT / SCHEDULED / CHECKED_IN / VERIFIED / IN_PROGRESS  →  CANCELLED      (COMPLETED & CANCELLED = status akhir)",
        size=14, bold=True, color=WARN, anchor=MSO_ANCHOR.MIDDLE)
table(s, ["Aturan", "Efek"], [
    ["Transisi ilegal", "Ditolak HTTP 422 — PUT /rencana-kirim/{id} tidak dapat mengubah status"],
    ["Edit item", "Hanya sah saat DRAFT/SCHEDULED; setelah CHECKED_IN item jadi acuan"],
    ["Gate In", "Hanya diterima saat SCHEDULED; duplikat ditolak"],
    ["Validasi selesai", "Otomatis membuat Receiving + ReceivingItems (status PENDING)"],
    ["Receiving selesai", "Dock kembali AVAILABLE, batch siap, email 'Proses Bongkar Selesai'"],
    ["Sesi Opname aktif", "Rencana tidak boleh diubah sampai sesi dibatalkan"],
], Inches(0.7), Inches(4.02), Inches(11.9), col_ratios=[1.1, 4], size=13, row_h=0.36)

# ══════════════════════════ 10. LANGKAH 1
s = new("Langkah 1 — Replenishment", "Peran: Admin Planner")
card(s, Inches(0.7), Inches(1.72), Inches(5.4), Inches(2.35), "Alur 5 Langkah", [
    "Buka **kalender** → pilih bulan & tanggal",
    "Isi **dokumen**: No. SJ/BA + Tanggal RKH",
    "Isi **kendaraan**: driver, ekspedisi, nopol",
    "Isi **pengirim & tujuan** (depo)",
    "Pilih **produk NPN** / import Excel → Simpan & Jadwal",
], PRIMARY, size=12.5, glyph="📋")
card(s, Inches(0.7), Inches(4.25), Inches(5.4), Inches(2.3), "Aturan & Keluaran", [
    "`po_transporter` **unik per depo** (duplikat 422)",
    "Kubikasi terhitung live dari data NPN",
    "Simpan → **email TO BLC, CC Cheker**",
    "Surat Jalan PDF (barcode Code39) dari browser",
], ACCENT, size=12.5, glyph="⚙️")
shot(s, "replenishment", Inches(6.5), Inches(1.85), Inches(6.1), "Layar Replenishment — kalender + daftar rencana", max_h=Inches(3.85))

# ══════════════════════════ 11. LANGKAH 2
s = new("Langkah 2 — Check In Point", "Peran: Security")
card(s, Inches(0.7), Inches(1.72), Inches(5.4), Inches(2.35), "Yang Dilakukan", [
    "Truk tiba → **scan barcode PO Transporter**",
    "Atau cari no. dokumen (pencarian cepat)",
    "Sistem cek status rencana → harus **SCHEDULED**",
    "Transaksi INBOUND tersimpan → **CHECKED_IN**",
], ORANGE, size=12.5, glyph="🚧")
card(s, Inches(0.7), Inches(4.25), Inches(5.4), Inches(2.3), "Validasi Backend", [
    "POST /api/v1/gate-transactions (permission `gate.create`)",
    "Duplikat check-in **ditolak 422**",
    "Statistik gate per depo & tanggal",
], PRIMARY, size=12.5, glyph="🔐")
shot(s, "gate", Inches(6.5), Inches(1.85), Inches(6.1), "Layar Check In Point — pencarian dokumen & scan", max_h=Inches(3.85))

# ══════════════════════════ 12. LANGKAH 3
s = new("Langkah 3 — Validasi Dokumen", "Peran: Cheker · titik paling kritis")
card(s, Inches(0.7), Inches(1.72), Inches(5.4), Inches(2.3), "Input per Item (scan)", [
    "`actual_qty` (wajib) vs expected",
    "`expired_date`, `barcode`, `notes`",
    "`condition`: GOOD / DAMAGED / MISSING",
    "`tier` & `tinggi_tumpukan` → Standarisasi (pending)",
], ACCENT, size=12.5, glyph="🔎")
card(s, Inches(0.7), Inches(4.15), Inches(5.4), Inches(2.4), "Otomatis oleh Sistem", [
    "1 row **batches** per produk → BATCH- & LBL-",
    "Idempoten — scan ulang tidak menggandakan batch",
    "Receiving + ReceivingItems dibuat (PENDING)",
    "Status rencana → **VERIFIED** · email terkirim",
], PRIMARY, size=12.5, glyph="🤖")
shot(s, "validasi-sj", Inches(6.5), Inches(1.85), Inches(6.1), "Layar Validasi Dokumen — daftar item & scan", max_h=Inches(3.85))

# ══════════════════════════ 13. LANGKAH 4
s = new("Langkah 4 — Check Out (Receiving)", "Peran: Cheker")
card(s, Inches(0.7), Inches(1.72), Inches(5.4), Inches(2.35), "Yang Dilakukan", [
    "Buka Check Out → **scan barcode PO**",
    "Barcode harus cocok dengan `po_transporter`",
    "Quick Scan juga dengan **nama driver**",
    "Klik **Selesai Bongkar** — tanpa scan ulang produk",
], ACCENT, size=12.5, glyph="📦")
card(s, Inches(0.7), Inches(4.25), Inches(5.4), Inches(2.3), "Dampak Sistem", [
    "Status → **COMPLETED** (status akhir)",
    "Dock kembali AVAILABLE · batch siap",
    "Email **Proses Bongkar Selesai** (TO BLC)",
    "Dasar laporan Inbound & Selisih Qty",
], PRIMARY, size=12.5, glyph="🏁")
shot(s, "receiving", Inches(6.5), Inches(1.85), Inches(6.1), "Layar Check Out — Quick Scan & progres bongkar", max_h=Inches(3.85))

# ══════════════════════════ 14. ALUR DATA & BARCODE
s = new("Alur Data & Barcode", "Kenapa cukup scan sekali")
flow(s, [("Daftar Barang RK", "rencana_kirim_items", "📋"),
         ("Snapshot Validasi", "sj_validation_items", "🔎"),
         ("Input Scan Item", "qty · expired · kondisi", "📲"),
         ("Batch Otomatis", "BATCH-XXXXXXXX", "🏷️"),
         ("Cetak Label", "QR + LBL-XXXXXXXX", "🖨️"),
         ("Receiving", "receiving_items", "📦")],
     Inches(1.7), height=Inches(1.3), label_size=12.5, sub_size=10)
card(s, Inches(0.7), Inches(3.45), Inches(3.85), Inches(3.05), "Titik Scan Tunggal", [
    "Produk **hanya discan** di Validasi SJ",
    "Hasil scan **diteruskan** ke Receiving, bukan diinput ulang",
    "Menghindari selisih akibat input dua kali",
    "Selisih qty terlihat langsung (expected vs actual)",
], ACCENT, size=12.5, glyph="🎯")
card(s, Inches(4.75), Inches(3.45), Inches(3.85), Inches(3.05), "Identitas Batch", [
    "`batch_number` BATCH-XXXXXXXX (unik)",
    "`barcode_label` LBL-XXXXXXXX",
    "QR code encode batch_number",
    "Teks: nama barang, qty, satuan",
    "ESC/POS via Web Bluetooth / dialog print",
], PRIMARY, size=12.5, glyph="🏷️")
card(s, Inches(8.8), Inches(3.45), Inches(3.8), Inches(3.05), "Dampak ke Prospek Data", [
    "Laporan **Selisih Qty** otomatis",
    "Monitoring **Batch Expired** per batch",
    "Kubikasi turunan dari NPN (tanpa simpan)",
    "Standarisasi tumpukan tercatat per produk",
], DARK, size=12.5, glyph="📈")

# ══════════════════════════ 15. OPNAME
s = new("Workflow Pendukung — Opname", "Batch Expired & Sesi Opname")
table(s, ["Status", "Sisa Hari", "Makna & Aksi"], [
    ["AMAN", "≥ 90", "Stok aman, FIFO normal"],
    ["PERHATIAN", "≥ 30", "Prioritaskan produk keluar"],
    ["KRITIS", "0 – 29", "Segera kirim / retur"],
    ["EXPIRED", "< 0", "Kedaluwarsa — wajib ditarik"],
], Inches(0.7), Inches(1.72), Inches(5.4), col_ratios=[1, 0.7, 2], size=12.5, row_h=0.38)
for i, (t, c) in enumerate([("AMAN", GREEN), ("PERHATIAN", ORANGE), ("KRITIS", WARN), ("EXPIRED", RGBColor(0x7F, 0x1D, 0x1D))]):
    pill(s, Inches(0.7) + i * Inches(1.38), Inches(3.75), Inches(1.25), Inches(0.34), t, c, fs=11)
card(s, Inches(0.7), Inches(4.3), Inches(5.4), Inches(2.25), "Generate Mingguan", [
    "Agregat ReceivingItem per depo + barang + expired",
    "Roll-forward batch lama yang belum EXPIRED",
    "Catat perubahan: BARU · TETAP · NAIK · TURUN",
], PRIMARY, size=12.5, glyph="🔄")
card(s, Inches(6.5), Inches(1.72), Inches(6.1), Inches(2.0), "Sesi Opname (Verifikasi Lapangan)", [
    "Buat sesi dari wave → **Start** → verifikasi item → **Complete**",
    "Wave terkunci selama sesi IN_PROGRESS",
    "Koreksi hanya qty_koreksi_crt / pcs",
], ACCENT, size=12.5, glyph="🧾")
shot(s, "opname-expired", Inches(6.5), Inches(4.05), Inches(6.1),
     "Layar Batch Expired — monitoring per minggu", max_h=Inches(2.45))

# ══════════════════════════ 16. BIAYA
s = new("Workflow Pendukung — Biaya THL & Retribusi", "Approval berjenjang")
flow(s, [("DRAFT", "pengisi", "✍️"), ("SUBMIT", "diajukan", "📤"), ("CHECKED", "supervisor", "🔍"),
         ("RLH_APPROVED", "RLH", "👍"), ("APPROVED", "RFAH/BLC/BAC", "💰")],
     Inches(1.7), height=Inches(1.05), label_size=14, sub_size=11)
rect(s, Inches(0.7), Inches(2.9), Inches(11.9), Inches(0.58), fill=SOFTRED, line=RGBColor(0xFE, 0xCA, 0xCA))
textbox(s, Inches(0.95), Inches(2.96), Inches(11.4), Inches(0.45),
        "REJECTED dapat terjadi di tahap mana pun   ·   REVISI mengembalikan dokumen ke pengisi",
        size=14, bold=True, color=WARN, anchor=MSO_ANCHOR.MIDDLE)
card(s, Inches(0.7), Inches(3.65), Inches(3.05), Inches(1.62), "Isi Dokumen", [
    "Detail biaya + tarif",
    "Foto dokumen JPEG ≤ 2 MB",
    "Import Excel massal",
    "Tarif retribusi per depo",
], PRIMARY, size=11.5, glyph="🧾")
card(s, Inches(3.9), Inches(3.65), Inches(3.05), Inches(1.62), "Mesin Approval", [
    "BiayaApprovalService",
    "ApprovalSectionResolver",
    "approval_assignments",
    "approval_histories (audit)",
], ACCENT, size=11.5, glyph="⚙️")
card(s, Inches(0.7), Inches(5.42), Inches(6.25), Inches(1.2), "Keluaran", [
    "PDF persetujuan · ekspor Excel per periode & depo",
    "Laporan Biaya THL / Retribusi · kandidat dari penerimaan berjalan",
], DARK, size=11.5, glyph="📤")
shot(s, "biaya-thl", Inches(7.25), Inches(3.7), Inches(5.35),
     "Layar Biaya THL — pipeline approval & detail biaya", max_h=Inches(2.5))

# ══════════════════════════ 17. LAPORAN & EMAIL
s = new("Workflow Pendukung — Laporan & Notifikasi", "Keluaran sistem")
textbox(s, Inches(0.7), Inches(1.65), Inches(5.4), Inches(0.32), "📊  7 HALAMAN LAPORAN", size=13, bold=True, color=ACCENT)
table(s, ["Laporan", "Sumber Data"], [
    ["Replenishment", "rencana_kirim + item"],
    ["Inbound", "penerimaan per depo"],
    ["Penerimaan & Expired", "gabungan penerimaan + batch"],
    ["Selisih Qty", "expected vs actual"],
    ["Biaya THL", "biaya_thl + detail"],
    ["Biaya Retribusi", "biaya_retribusi + detail"],
    ["Batch Expired", "opname_expired_waves"],
], Inches(0.7), Inches(2.05), Inches(5.4), col_ratios=[1.3, 1.6], size=12, row_h=0.34)
textbox(s, Inches(0.7), Inches(5.35), Inches(5.4), Inches(0.8),
        "Filter tanggal + depo → ekspor Excel → gate permission reports.view.",
        size=12.5, color=MUTED)
shot(s, "reports-selisih", Inches(6.5), Inches(1.8), Inches(6.1),
     "Laporan Selisih Qty — filter periode + ekspor", max_h=Inches(3.0))
card(s, Inches(6.5), Inches(5.3), Inches(6.1), Inches(1.5), "Notifikasi Email Otomatis", [
    "rencana_kirim/created → TO BLC, CC Cheker",
    "receiving/completed → TO BLC, CC Admin Distribusi",
    "sj_ba/validated & selisih_qty → penerima khusus",
], PRIMARY, size=11.5, glyph="✉️")

# ══════════════════════════ 18. KEAMANAN
s = new("Keamanan & Otorisasi", "Siapa boleh melakukan apa")
table(s, ["Aspek", "Implementasi"], [
    ["Autentikasi Web", "Sanctum stateful cookie — tanpa Bearer token di browser"],
    ["CSRF", "withCredentials + withXSRFToken, exchange /sanctum/csrf-cookie"],
    ["Proxy", "proxy.ts meneruskan /api/* & /sanctum/* (cookie terikat domain)"],
    ["Otorisasi", "Middleware permission:* → CheckPermission → 403 + daftar permission"],
    ["Audit", "TrackUserActivity + queue LogUserActivity → user_activity_logs"],
    ["Rate Limit", "POST /login throttle 10 kali per menit"],
    ["Integritas Status", "Transisi status dikunci backend"],
    ["Reset Data", "database.purge hanya admin_master"],
], Inches(0.7), Inches(1.72), Inches(7.3), col_ratios=[1.05, 3.4], size=12.5, row_h=0.4)
card(s, Inches(8.25), Inches(1.72), Inches(4.35), Inches(1.7), "14 Role", [
    "admin_master · staff_logistic · rlh · rfah · rh",
    "blc · bac · admin_distirbusi · admin_planner",
    "cheker · sac · operator · ga · bh",
], ACCENT, size=11.5, glyph="🛡️")
shot(s, "roles", Inches(8.25), Inches(3.7), Inches(4.35), "Layar Roles & Permissions", max_h=Inches(2.7))

# ══════════════════════════ 19. PERAN
s = new("Peran & Tanggung Jawab", "RACI pada workflow utama")
table(s, ["Peran", "Tahap", "Tanggung Jawab"], [
    ["Admin Planner", "1. Replenishment", "Rencana kirim, dokumen, kendaraan, daftar barang"],
    ["Security", "2. Check In", "Gate-in kendaraan, pencatatan transaksi masuk"],
    ["Cheker", "3. Validasi SJ", "Cocokkan dokumen vs fisik, qty/kondisi/expired, cetak label"],
    ["Cheker", "4. Check Out", "Bongkar sampai selesai, status COMPLETED"],
    ["Admin Distribusi", "4. Check Out", "Menerima notifikasi & monitoring distribusi"],
    ["RLH / RFAH / BLC / BAC", "Biaya", "Approval berjenjang biaya THL & retribusi"],
    ["Admin Master", "Sistem", "Users, roles, approval, email, reset data"],
    ["RH / GA / SAC / Operator", "Pendukung", "Viewer & pelaksana operasional depot"],
], Inches(0.7), Inches(1.72), Inches(11.9), col_ratios=[1.35, 1.05, 4], size=13, row_h=0.5)

# ══════════════════════════ 20. MANFAAT
s = new("Manfaat & Indikator Keberhasilan", "Dampak yang diharapkan")
items = [
    ("Akurasi Stok", "📊", "Qty aktual tercatat sekali saat validasi — selisih terdeteksi dini.", PRIMARY),
    ("Barang Tidak Salah Terima", "🛡️", "Bongkar hanya setelah SJ/BA tercocokkan dengan fisik.", ACCENT),
    ("Risiko Expired Turun", "⏳", "Peringatan AMAN/PERHATIAN/KRITIS/EXPIRED + verifikasi lapangan.", WARN),
    ("Biaya Terkendali", "💰", "Biaya cair hanya setelah approval berjenjang & jejak audit.", ACCENT),
    ("Kerja Kertas Minim", "📄", "Surat Jalan, label, dan laporan dihasilkan sistem.", PRIMARY),
    ("Jejak Audit Penuh", "🔍", "Status, batch, aktivitas user, riwayat approval tersimpan.", DARK),
]
for i, (h, g, d, c) in enumerate(items):
    left = Inches(0.7) + (i % 3) * Inches(4.05)
    top = Inches(1.75) + (i // 3) * Inches(2.4)
    rect(s, left, top, Inches(3.8), Inches(2.15), fill=CARD, line=BORDER)
    rect(s, left, top, Inches(3.8), Pt(5), fill=c)
    icon_badge(s, left + Inches(0.2), top + Inches(0.25), Inches(0.42), g, c, fs=17)
    textbox(s, left + Inches(0.75), top + Inches(0.28), Inches(2.9), Inches(0.6), h, size=15.5, bold=True, color=c)
    textbox(s, left + Inches(0.2), top + Inches(0.95), Inches(3.4), Inches(1.1), d, size=13, color=TEXT, spacing=1.1)

# ══════════════════════════ 21. TEKNIS
s = new("Catatan Teknis & Batasan", "Yang perlu diketahui tim")
card(s, Inches(0.7), Inches(1.72), Inches(5.8), Inches(2.3), "Fakta Teknis", [
    "Tanpa `any` di seluruh kode frontend TypeScript",
    "Tanpa `setInterval` — tidak ada polling",
    "Auth state hanya di memori (Zustand tanpa persist)",
    "PDF Surat Jalan dirender client-side (jsPDF)",
], PRIMARY, size=12.5, glyph="🛠️")
card(s, Inches(6.8), Inches(1.72), Inches(5.8), Inches(2.3), "Pengujian", [
    "Backend: PHPUnit — 99 file test, SQLite",
    "Frontend: Vitest 5 + testing-library (coverage 30%)",
    "Alur end-to-end diuji langsung ke API /api/v1",
    "Rujukan: docs/QA_REPORT.md",
], ACCENT, size=12.5, glyph="🧪")
card(s, Inches(0.7), Inches(4.22), Inches(5.8), Inches(2.3), "Modul yang BELUM Ada", [
    "Appointment, ASN, Yard, Gates Master",
    "Receiving Schedule, Put Away",
    "Opname non-expired, FEFO",
], WARN, size=12.5, glyph="⚠️")
card(s, Inches(6.8), Inches(4.22), Inches(5.8), Inches(2.3), "Dokumen Terkait", [
    "docs/project_flow.md — struktur & alur lengkap",
    "docs/THL_RETRIBUSI_SYSTEM_FLOW.md — detail biaya",
    "docs/WORKFLOW_TUJUAN_SISTEM.md — versi teks deck ini",
], PRIMARY, size=12.5, glyph="📚")

# ══════════════════════════ 22. PENUTUP
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=DARK)
rect(s, 0, 0, Inches(0.35), H, fill=ACCENT)
textbox(s, Inches(1.1), Inches(1.5), Inches(11.0), Inches(0.4), "RINGKASAN", size=15, bold=True, color=ACCENT)
textbox(s, Inches(1.1), Inches(2.0), Inches(11.2), Inches(1.2),
        "Satu alur barang, satu sumber kebenaran", size=40, bold=True, color=WHITE)
rect(s, Inches(1.12), Inches(3.2), Inches(1.8), Pt(5), fill=ACCENT)
bullets(s, [
    "Replenishment → Check In → Validasi SJ → Check Out: satu jalur wajib dengan status terkunci.",
    "Produk discan sekali di Validasi SJ; batch, label, dan receiving dibuat otomatis.",
    "Opname expired & sesi opname menjaga stok tetap layak jual.",
    "Biaya THL & retribusi cair hanya setelah approval berjenjang.",
    "RBAC + audit trail membuat setiap aksi dapat dipertanggungjawabkan.",
], Inches(1.12), Inches(3.65), Inches(10.6), Inches(2.4), size=17, gap=13, color=RGBColor(0xE0, 0xE7, 0xFF))
try:
    s.shapes.add_picture(LOGO, Inches(10.7), Inches(6.2), width=Inches(1.9))
except Exception:
    pass
textbox(s, Inches(1.1), Inches(6.4), Inches(9.0), Inches(0.5),
        "ONEZENTRA Warehouse System · Next.js 16 + Laravel 11 · docs/ONEZENTRA_WORKFLOW.pptx",
        size=12, color=RGBColor(0xA5, 0xB4, 0xFC))
_page["n"] += 1

prs.save(OUT)
print("saved:", OUT)
